"""Convert uploaded documents and images into CSV tables for proof analysis."""

from __future__ import annotations

import base64
import re
from io import BytesIO
from pathlib import Path

import pandas as pd

SUPPORTED_EXTENSIONS = {".csv", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp"}
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_IMAGE_COUNT = 3


def _safe_name(filename: str, fallback: str) -> str:
    stem = Path(filename).stem
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", stem).strip("_") or fallback
    return safe[:80]


def _image_to_text(data: bytes, mime_type: str, api_key: str, provider: str) -> str:
    from agent.llm_prompt import _resolve_model

    prompt = (
        "Read this image, including any chart labels and visible values. Return a concise, "
        "literal transcription and describe any chart data as labeled values. Clearly say "
        "when a value is approximate or unreadable; never invent missing values."
    )
    if provider == "gemini":
        from google import genai
        from google.genai import types

        client = genai.Client(api_key=api_key)
        result = client.models.generate_content(
            model=_resolve_model(provider="gemini"),
            contents=[types.Part.from_bytes(data=data, mime_type=mime_type), prompt],
            config={"temperature": 0},
        )
        return (getattr(result, "text", None) or "").strip()

    from groq import Groq

    encoded = base64.b64encode(data).decode("ascii")
    result = Groq(api_key=api_key).chat.completions.create(
        model="qwen/qwen3.8-27b",
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{encoded}"}},
                ],
            }
        ],
        temperature=0,
        max_completion_tokens=2048,
    )
    return (result.choices[0].message.content or "").strip()


def stage_uploads(
    uploads: dict[str, bytes],
    destination: Path,
    *,
    api_key: str | None = None,
    provider: str | None = None,
) -> list[str]:
    """Stage CSVs and turn PDF/DOCX/image uploads into proof-readable CSVs."""
    destination.mkdir(parents=True, exist_ok=True)
    from agent.llm_prompt import _resolve_api_key, _resolve_provider

    resolved_provider = (provider or _resolve_provider()).lower()
    key = _resolve_api_key(api_key, resolved_provider)
    doc_rows: list[dict[str, str]] = []
    image_count = 0
    staged: list[str] = []
    used_names: set[str] = set()

    for original_name, data in uploads.items():
        ext = Path(original_name).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {Path(original_name).name}")
        if len(data) > MAX_UPLOAD_BYTES:
            raise ValueError(f"{Path(original_name).name} exceeds the 20 MB upload limit.")
        stem = _safe_name(original_name, f"upload_{len(staged) + 1}")
        name = stem
        suffix = 2
        while name.lower() in used_names:
            name = f"{stem}_{suffix}"
            suffix += 1
        used_names.add(name.lower())

        if ext == ".csv":
            filename = f"{name}.csv"
            (destination / filename).write_bytes(data)
            # Validate early and give a useful upload error.
            pd.read_csv(destination / filename)
            staged.append(filename)
        elif ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            try:
                import fitz

                rendered_pdf = fitz.open(stream=data, filetype="pdf")
            except ImportError:
                fitz = None
                rendered_pdf = None
            for page_number, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if rendered_pdf is not None and (len(text) < 120 or page.images):
                    if not key:
                        raise ValueError("An API key is required to read scanned or chart pages in a PDF.")
                    rendered = rendered_pdf[page_number - 1].get_pixmap(
                        matrix=fitz.Matrix(1.5, 1.5), alpha=False
                    )
                    visual = _image_to_text(rendered.tobytes("png"), "image/png", key, resolved_provider)
                    if visual:
                        text = f"{text}\n[Visual page reading]\n{visual}".strip()
                doc_rows.append({
                    "source": Path(original_name).name,
                    "kind": "PDF text and visual reading",
                    "location": str(page_number),
                    "content": text,
                })
        elif ext == ".docx":
            from docx import Document

            document = Document(BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
            for table_number, table in enumerate(document.tables, start=1):
                table_text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
                text += f"\n[Table {table_number}]\n{table_text}"
            doc_rows.append({"source": Path(original_name).name, "kind": "Word text", "location": "document", "content": text.strip()})
        else:
            image_count += 1
            if image_count > MAX_IMAGE_COUNT:
                raise ValueError("Upload up to 3 images at a time.")
            if not key:
                raise ValueError("An API key is required to read graph images. Configure GROQ_API_KEY or GEMINI_API_KEY.")
            mime = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}[ext]
            extracted = _image_to_text(data, mime, key, resolved_provider)
            doc_rows.append({"source": Path(original_name).name, "kind": "Image/chart extraction", "location": "image", "content": extracted})

    if doc_rows:
        filename = "uploaded_documents.csv"
        if Path(filename).stem.casefold() in used_names or (destination / filename).exists():
            filename = "extracted_documents.csv"
        counter = 2
        while (destination / filename).exists():
            filename = f"extracted_documents_{counter}.csv"
            counter += 1
        pd.DataFrame(doc_rows, columns=["source", "kind", "location", "content"]).to_csv(
            destination / filename, index=False
        )
        staged.append(filename)
    return staged
